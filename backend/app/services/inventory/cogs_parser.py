"""Parse the weekly AFS inventory-count spreadsheet.

Layout (one sheet):
  * A banner + a few label rows (Store / Date).
  * Category sections. A section header row has the word "TOTAL" in the last
    column (H). Under it, each priced line has: name (A), pack size (B),
    unit (C), cases (D), packs (E), weight (F), case $ (G), extended total (H).
    Sub-labels (Frozen, Dry, Sauces…) carry no case price and are skipped as
    rows but keep the current category.
  * A trailing summary table beginning at a "CATEGORY" row with per-category
    "This Week's Purchases" the user may have filled in.

The parser is defensive: it keys off cell *shape* (a real item always has a
numeric case price in column G), not fixed row numbers, so added/removed items
don't break it.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal, InvalidOperation

import openpyxl

from app.db.models.inventory import INVENTORY_CATEGORIES

# Section label (uppercased) → canonical category. First substring hit wins.
_SECTION_MAP = (
    ("PROTEIN", "Protein"),
    ("CHEESECAKE", "Cheesecake"),
    ("GROCERY", "Grocery"),
    ("PRODUCE", "Produce"),
    ("BEVERAGE", "Beverage"),
    ("PACKAGING", "Packaging"),
)
_CANON = set(INVENTORY_CATEGORIES)


def _dec(v) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _canon_category(label: str) -> str | None:
    up = label.upper()
    for needle, canon in _SECTION_MAP:
        if needle in up:
            return canon
    return None


class ParsedCount:
    def __init__(self) -> None:
        self.week_ending: date | None = None
        self.items: list[dict] = []
        # canonical category -> Decimal purchases entered on the summary table
        self.purchases: dict[str, Decimal] = {}

    @property
    def category_ending(self) -> dict[str, Decimal]:
        totals: dict[str, Decimal] = {c: Decimal("0") for c in INVENTORY_CATEGORIES}
        for it in self.items:
            totals[it["category"]] += it["extended_total"]
        return totals


def parse_cogs_sheet(file_bytes: bytes) -> ParsedCount:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = ParsedCount()

    current: str | None = None
    in_summary = False
    sort = 0

    for row in ws.iter_rows(values_only=True):
        cells = list(row) + [None] * (8 - len(row)) if len(row) < 8 else list(row)
        a = cells[0]
        a_str = str(a).strip() if a is not None else ""
        last = cells[7]  # column H

        # Date label → capture the adjacent value if it's a real date.
        if a_str.upper().startswith("DATE") and out.week_ending is None:
            for v in cells[1:]:
                if isinstance(v, date):
                    out.week_ending = v if not hasattr(v, "date") else v.date()
                    break

        # Enter the trailing summary table.
        if a_str.upper() == "CATEGORY":
            in_summary = True
            continue

        if in_summary:
            canon = a_str if a_str in _CANON else None
            if canon:
                # Column E (index 4) = "This Week's Purchases".
                p = _dec(cells[4])
                if p is not None:
                    out.purchases[canon] = p
            continue

        # Section header: "TOTAL" sits in the last column.
        if isinstance(last, str) and last.strip().upper() == "TOTAL":
            mapped = _canon_category(a_str)
            if mapped:
                current = mapped
            continue

        # A priced line item: has a case price in column G (index 6).
        price = _dec(cells[6])
        if current and a_str and price is not None:
            ext = _dec(cells[7]) or Decimal("0")
            out.items.append({
                "category": current,
                "item_name": a_str[:256],
                "pack_size": (str(cells[1]).strip()[:64] if cells[1] not in (None, "") else None),
                "unit": (str(cells[2]).strip()[:16] if cells[2] not in (None, "") else None),
                "cases": _dec(cells[3]),
                "packs": _dec(cells[4]),
                "weight": _dec(cells[5]),
                "unit_price": price,
                "extended_total": ext,
                "sort_order": sort,
            })
            sort += 1

    return out
